"""Rebuild V3 tables, charts and Chinese workbook from public evidence only."""
from __future__ import annotations

import json
import statistics
from collections import defaultdict
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "evidence/v3"
LABEL = {"upgraded": "V2", "research_full": "V3 full", "research": "V3 budget"}


def load(path):
    return json.loads(path.read_text())


def sheet(wb, name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([json.dumps(x, ensure_ascii=False) if isinstance(x, (dict, list)) else x
                   for x in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="173D75")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(70, max(16, width + 2))
        for cell in col[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def main():
    rows = [load(p) for p in sorted((OUT / "trials").glob("*.json"))]
    held = [r for r in rows if r["case"]["split"] == "heldout"]
    groups = defaultdict(list)
    for row in held:
        suite = "context" if row["case"]["mode"] == "context" else "tasks"
        groups[(suite, row["arm"])].append(row)
    summaries = []
    for (suite, arm), part in sorted(groups.items()):
        usage = [r["usage"] for r in part]
        summaries.append({
            "suite": suite, "arm": arm, "n": len(part), "passed": sum(r["passed"] for r in part),
            "mean_total_tokens": statistics.mean(u["total_tokens"] for u in usage),
            "mean_input_tokens": statistics.mean(u["input_tokens"] for u in usage),
            "mean_output_tokens": statistics.mean(u["output_tokens"] for u in usage),
            "mean_first_request_input_tokens": statistics.mean(
                r["calls"][0]["usage"]["input_tokens"] for r in part),
            "mean_seconds": statistics.mean(r["elapsed_seconds"] for r in part),
            "median_seconds": statistics.median(r["elapsed_seconds"] for r in part),
            "api_attempts": sum(u["api_attempts"] for u in usage),
            "usage_complete": all(u["usage_complete"] for u in usage),
            "observed_cached_tokens": sum(u["cached_tokens"] for u in usage),
            "tool_errors": sum(e.get("is_error", False) for r in part for e in r["events"]
                               if e.get("event_kind") == "tool_end"),
        })
    study = load(ROOT / "evidence/study/result.json")
    scheduler = load(OUT / "scheduler.json")
    scheduling = []
    for mode in ("serial", "parallel", "barrier"):
        part = [r for r in scheduler["trials"] if r["mode"] == mode]
        reads = [r["elapsed_seconds"] for r in part if r["case"] == "independent_reads"]
        scheduling.append({"mode": mode, "passed": sum(r["contract_pass"] for r in part),
                           "n": len(part), "six_reads_mean_seconds": statistics.mean(reads)})
    summary = {"basis": "Frozen project engineering criteria; not human gold",
               "heldout": summaries, "scheduler": scheduling,
               "study": study["summary"], "public_trials_including_dev": len(rows)}
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
    context = {r["arm"]: r for r in summaries if r["suite"] == "context"}
    reduction = 1 - context["research"]["mean_total_tokens"] / context["research_full"]["mean_total_tokens"]
    relative_v2 = 1 - context["research"]["mean_total_tokens"] / context["upgraded"]["mean_total_tokens"]
    # Exportable scientific figure. Synthetic tasks and measured metric named explicitly.
    fig, axes = plt.subplots(1, 2, figsize=(10, 4), constrained_layout=True)
    arms = ("upgraded", "research_full", "research")
    values = [context[a]["mean_total_tokens"] for a in arms]
    bars = axes[0].bar([LABEL[a] for a in arms], values, color=["#8e9fbe", "#c4cddb", "#216649"])
    axes[0].bar_label(bars, fmt="%.0f", padding=3)
    axes[0].set(ylabel="Mean total provider tokens / trial", ylim=(0, max(values) * 1.18),
                title="Synthetic heldout context · 32 trials / arm")
    axes[0].spines[["top", "right"]].set_visible(False)
    for i, method in enumerate(("pixels_svm", "pca_svm")):
        data = [next(r for r in study["summary"] if r["method"] == method and r["condition"] == c)
                for c in ("clean", "noise", "occlusion")]
        x = [j + (i - 0.5) * 0.35 for j in range(3)]
        axes[1].bar(x, [r["mean_accuracy"] * 100 for r in data], width=0.35,
                    yerr=[r["std_accuracy"] * 100 for r in data], capsize=3,
                    color=["#8e9fbe", "#216649"][i], label=method)
    axes[1].set(xticks=[0, 1, 2], xticklabels=["Clean", "Noise σ=3", "Occlusion 2×2"],
                ylabel="Test accuracy (%)", ylim=(0, 110), title="Real digits study · 3 seeds · mean ± SD")
    axes[1].legend(frameon=False, fontsize=8)
    axes[1].spines[["top", "right"]].set_visible(False)
    fig.savefig(OUT / "results.png", dpi=180)
    fig.savefig(OUT / "results.svg")
    plt.close(fig)
    svg = OUT / "results.svg"
    svg.write_text("\n".join(line.rstrip() for line in svg.read_text().splitlines()) + "\n")
    audit = load(OUT / "semantic-audit.json")
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "阅读口径", ["项目", "说明"], [
        ["输入", "冻结协议、相同模型和工具、128次留出调用任务；另保留36次开发试验"],
        ["处理", "独立循环、完整轮次上下文、按需工件化、读并行写屏障"],
        ["输出", "逐题trace、请求/返回、完整usage、精确标识判据、实测预测与区间"],
        ["金标来源", "工程规则和程序化随机标识；没有真实人评金标"],
        ["原创边界", "独立工程实现；复用DTO、传输、V2决策账本和标准机器学习算法"],
        ["正常任务失败", "V2/V3各14/16；4次研究回答未附完整方案ID，任务本身未明确要求该字段"],
        ["上下文边界", "16个合成会话×2次；共享开发题生成机制，不是外部benchmark"],
        ["Token", "供应商全部可观测调用用量，不等同价格；字符预算不是token上限"],
        ["模型", "doubao-seed-2-0-lite-260428；temperature=0；thinking disabled"],
        ["独立性", "AI语义复核由开发助手完成，不等同独立盲评"],
        ["研究边界", "digits、3seed测试集有重叠；区间条件于固定拆分和已训练模型"],
        ["负结果", "噪声准确率接近随机；遮挡下PCA平均更差；未证明总体鲁棒性收益"],
    ])
    sheet(wb, "留出汇总", ["任务集", "版本", "样本次", "通过", "平均总Token", "平均输入Token",
          "平均输出Token", "首请求平均输入Token", "平均秒", "中位秒", "API次数", "用量完整", "工具错误"],
          [[r[k] for k in ("suite", "arm", "n", "passed", "mean_total_tokens", "mean_input_tokens",
                           "mean_output_tokens", "mean_first_request_input_tokens", "mean_seconds",
                           "median_seconds", "api_attempts", "usage_complete", "tool_errors")] for r in summaries])
    sheet(wb, "全部逐题", ["试验ID", "划分", "版本", "通过", "未通过判据", "总Token", "API次数", "秒", "错误", "证据路径"],
          [[r["id"], r["case"]["split"], r["arm"], r["passed"],
            [k for k, v in r["checks"].items() if not v], r["usage"]["total_tokens"],
            r["usage"]["api_attempts"], r["elapsed_seconds"], r["errors"], f"trials/{r['id']}.json"]
           for r in rows])
    sheet(wb, "语义问题", ["试验ID", "问题类型", "原始证据", "边界或后续"],
          [[r["trials"], r["type"], r["evidence"], r.get("limitation", r.get("recommended_followup", ""))]
           for r in audit["findings"]])
    sheet(wb, "调度消融", ["模式", "通过", "样本次", "六次读调用均秒"],
          [[r["mode"], r["passed"], r["n"], r["six_reads_mean_seconds"]] for r in scheduling])
    sheet(wb, "研究均值", ["条件", "方法", "平均准确率", "跨种子标准差"],
          [[r["condition"], r["method"], r["mean_accuracy"], r["std_accuracy"]] for r in study["summary"]])
    sheet(wb, "配对区间", ["种子", "条件", "PCA减像素", "95%下界", "95%上界", "Bootstrap次数"],
          [[r[k] for k in ("seed", "condition", "pca_minus_pixels", "paired_bootstrap_95_low",
                           "paired_bootstrap_95_high", "bootstrap_samples")] for r in study["paired_comparisons"]])
    sheet(wb, "完整题目", ["试验ID", "用户Prompt", "系统Prompt", "工具Schema"],
          [[r["id"], r["case"]["prompt"], r["system_prompt"], r["tool_schemas"]] for r in held])
    wb.save(OUT / "原创机制与实验证据.xlsx")
    lines = [
        "# V3 留出实验与实战报告", "",
        "V3 的主要可量化收益是上下文用量与工具执行顺序的可解释性。当前结果没有证明正常任务成功率提高。",
        "全部数字由 `python -m scripts.report_research` 从公开 JSON 重建，失败不剔除。", "",
        "## 冻结与样本", "",
        "开发36次、留出128次，合计164次真实模型试验。正常任务8题×2版本×2次；"
        "合成上下文16会话×3版本×2次。模型为 doubao-seed-2-0-lite-260428，temperature=0、关闭thinking。"
        "每个批次并发4；正常任务和上下文批次曾同时运行，跨批次最高并发8。"
        "题目为项目自建，两个重复共享题目；不存在人评金标或开放领域泛化准确率。", "",
        "`benchmarks/v3/protocol.json` 在实现前冻结；完整上下文生成器和后端源码在开发后、留出前冻结。"
        "公开 export-manifest 记录原始与脱敏版本 SHA256，并验证冻结源码未变。", "",
        "## 留出结果", "",
        "| 任务 | 版本 | 严格通过 | 平均总Token | 平均秒 | API次数 | 工具错误 |",
        "|---|---|---:|---:|---:|---:|---:|",
    ]
    lines += [f"| {r['suite']} | {LABEL[r['arm']]} | {r['passed']}/{r['n']} | "
              f"{r['mean_total_tokens']:,.1f} | {r['mean_seconds']:.2f} | {r['api_attempts']} | {r['tool_errors']} |"
              for r in summaries]
    lines += [
        "", f"上下文总Token：V3预算版相对同引擎完整历史下降 **{reduction:.1%}**，相对V2下降 **{relative_v2:.1%}**。"
        "96次均命中精确随机标识，所有128次留出usage记录完整。", "",
        "总Token包含错误调用和重读。完整历史组在归档题额外尝试了不存在的工件ID，"
        "V2也出现无效工件读取；这些恢复成本放大了总用量差。故74.4%不能表述成所有场景的纯压缩率。"
        "summary.json另给首请求输入用量。用户修正子集上V2平均3956.5 token，"
        "V3平均4417.1；V3并非所有子场景更省。工具错误不等于API错误，也不一定导致最终失败。", "",
        "V3预算上限18000指包含系统、Schema与消息的canonical DTO JSON字符数；不是真实token或wire bytes。"
        "预算必要内容过大时明确失败。全历史组无预算约束。V3无LLM摘要调用。", "",
        "## 失败与语义复核", "",
        "V2/V3各有两次研究任务未在回答附上完整方案哈希，严格结果均14/16。"
        "工具执行、六组均值均正确；该题Prompt未明确要求完整ID，而冻结grader要求它。"
        "报告保留原分，公开这一评测限制，不改写为16/16。", "",
        "AI复核还发现：四次ViT答案出现原文没有的异常token；一次V2答案误称遮挡区间上界全部≤0；"
        "一次V3拒答虽结论合适，却把DETR移除FFN的消融误称为纯Transformer。"
        "引用门禁只证明来源存在且本轮已读取，不证明命题被原文支持。"
        "详见 [逐条语义问题](../evidence/v3/semantic-audit.json)，这些结果不是独立人评准确率。", "",
        "## 调度消融（人工延迟）", "",
        "| 模式 | 工程判据 | 六次只读调用均秒 |", "|---|---:|---:|",
    ]
    lines += [f"| {r['mode']} | {r['passed']}/{r['n']} | {r['six_reads_mean_seconds']:.4f} |" for r in scheduling]
    lines += [
        "", "每模式8种场景×5次，相同调用清单。直接并行在write→read的5次回放中均读到旧值。"
        "屏障模式40/40，连续只读仍可并行。延迟由脚本注入，不能当生产性能或任意副作用保证。", "",
        "## 真实研究", "", "| 条件 | 方法 | 准确率均值 | 跨seed标准差 |", "|---|---|---:|---:|",
    ]
    lines += [f"| {r['condition']} | {r['method']} | {r['mean_accuracy']:.2%} | {r['std_accuracy']:.2%} |"
              for r in study["summary"]]
    lines += [
        "", "固定seed17/43/89、60/20/20分割、PCA24、C=0.1/1/10、噪声σ=3、2×2遮挡。"
        "变换仅拟合训练集，C仅由clean验证集选择，不重拟合验证集。"
        "每seed两个方法共享测试样本和干扰，区间为2000次配对bootstrap。"
        "不同seed测试集重叠，不能合并为独立样本；区间条件于固定拆分与训练模型。", "",
        "两个方法在噪声下均接近随机分类；PCA在遮挡下平均更差。这些负结果保留，"
        "不宣称PCA改善整体鲁棒性，也不宣称复现大型CV论文。"
        "独立验证脚本已重算18组指标、所有均值/标准差与区间，并核对数据集、拆分、干扰和方案哈希。", "",
        "## 实战与复验", "",
        "桌面与手机各完成：固定方案→执行或同方案复用→6行数值匹配→真实下载并校验SHA→"
        "V3调用read_study解释→Markdown导出。2项浏览器用例通过，未出现页面JS错误或横向溢出。"
        "这两项使用已有默认方案时验证的是复用；首次UI真实执行另由本地记录和独立指标核验支持。", "",
        "UI通过指动作、数值表、下载与导出通过。桌面模型回答仍把noise区间下界说成全部>0，"
        "并把occlusion区间上界说成全部≤0；已在demo与语义审计中标注并保留，未重试筛选截图。", "",
        "![实验结果](../evidence/v3/results.png)", "",
        "![真实研究界面](../evidence/demo/desktop-study.png)", "",
        "```bash", "python -m scripts.verify_study evidence/study",
        "python -m scripts.reproduce_study", "python -m scripts.report_research", "```", "",
        "[中文Excel](../evidence/v3/原创机制与实验证据.xlsx) · "
        "[逐题请求与结果](../evidence/v3/trials/) · [公开实战记录](../evidence/demo/) · "
        "[开发失败](v3-development.md)", "",
    ]
    (ROOT / "docs/research-evaluation.md").write_text("\n".join(lines))
    # Self-contained full prompts and role-specific tool schemas.
    prompts = ["# V3 完整 Prompt 与工具 Schema", "",
               "以下取自留出试验实际请求，模型transport ID已脱敏。每题原文也保留在protocol.json。",
               "合成历史由 scripts.benchmark_research.context_cases 生成。", ""]
    seen = set()
    for r in held:
        role = r["case"]["mode"]
        if role in seen:
            continue
        seen.add(role)
        prompts += [f"## {role} 系统 Prompt", "", "```text", r["system_prompt"], "```", "",
                    "```json", json.dumps(r["tool_schemas"], ensure_ascii=False, indent=2), "```", ""]
    prompts += ["## 每题用户 Prompt", ""]
    seen.clear()
    for r in held:
        if r["case"]["id"] not in seen:
            seen.add(r["case"]["id"])
            prompts += [f"### {r['case']['id']}", "", "```text", r["case"]["prompt"], "```", ""]
    (ROOT / "docs/research-prompts.md").write_text("\n".join(prompts))
    print(f"Built report, workbook, prompts and plots from {len(rows)} public trials")


if __name__ == "__main__":
    main()
