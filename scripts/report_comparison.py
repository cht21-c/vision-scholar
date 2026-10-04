"""Regrade saved traces, replay citation gates, export evidence-linked reports.

No model calls, no reruns, no mutation of original trials. Provider usage is
recomputed per attempt so missing summary-call usage cannot masquerade as zero.
"""
from __future__ import annotations

import collections
import json
import re
import statistics

import matplotlib
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.experiments import ExperimentInput
from backend.store import Store
from backend.tools import (
    Arguments,
    ArxivImportInput,
    ArxivSearchInput,
    CodeInput,
    LogInput,
    NoteInput,
    ReadInput,
    SearchInput,
)
from backend.upgrades import verify_citations_with_ranges
from scripts.compare_harnesses import OUTPUT, write_json

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

NAMES = {"legacy": "V0·LangGraph重建", "openharness": "V1·OpenHarness", "upgraded": "V2·应用升级"}
GROUPS = {"normal": "正常任务", "context": "合成上下文", "canary": "连通小样"}
CASES = {
    "resnet": "ResNet残差", "vit": "ViT分块", "clip": "CLIP目标", "detr": "DETR匹配",
    "compare": "跨论文比较", "unknown": "未知信息拒答", "code": "源码定位",
    "digits": "真实digits", "log": "合成日志", "memory": "保存与回读",
    "short": "短历史", "tool_bloat": "冗长工具结果", "buried_correction": "埋藏修正",
    "full_summary": "完整摘要压力",
}
SCHEMAS = dict(search_papers=SearchInput, read_paper=ReadInput, inspect_code=CodeInput,
               run_experiment=ExperimentInput, analyze_log=LogInput, save_note=NoteInput,
               list_notes=Arguments, search_arxiv=ArxivSearchInput, import_arxiv=ArxivImportInput)
CORRECTIONS = [
    {"id": "numeric-format-v2", "reason": "digits接受0.9778或97.78%的等价形式，原百分数判据漏报。",
     "scope": "同一原始回答；未重跑；实际工具结果作数值依据。"},
    {"id": "uncertainty-wording-v2",
     "reason": "拒答增加无法…确认等语序；无法根据这四篇论文确认也是明确拒答。",
     "scope": "只修辞面标记；不表示附带论文结论语义正确，语义另列AI复核。"},
]


def events(turn, kind):
    return [e for e in turn["events"] if e["type"] == kind]


def evidence_from(turn):
    evidence = dict(turn.get("evidence") or {})
    for e in events(turn, "ToolExecutionCompleted"):
        if e["is_error"]:
            continue
        try:
            value = json.loads(e["output"])
        except ValueError:
            continue
        if e["tool_name"] == "inspect_code":
            evidence[value["id"]] = value
        elif e["tool_name"] in {"search_papers", "read_paper"}:
            for item in value.get("results", value.get("chunks", [])):
                evidence[item["id"]] = item
    return evidence


def usage_of(calls):
    out = {"attempts": len(calls), "usage_records": sum(c["usage"] is not None for c in calls)}
    out["usage_complete"] = out["attempts"] == out["usage_records"]
    for key in ("input_tokens", "output_tokens", "total_tokens", "cached_tokens", "reasoning_tokens"):
        values = [c["usage"].get(key) for c in calls if c.get("usage") is not None]
        known = [v for v in values if v is not None]
        observed = sum(known) if known else None
        out["observed_" + key] = observed
        out[key] = observed if out["usage_complete"] and all(v is not None for v in values) else None
    return out


def regrade(trial):
    checks = dict(trial["checks"])
    numeric_only = dict(checks)
    notes = []
    answer = "\n".join(t["answer"] for t in trial["turns"])
    if trial["case_id"] == "digits":
        results = [json.loads(e["output"])["result"]["results"]
                   for t in trial["turns"] for e in events(t, "ToolExecutionCompleted")
                   if e["tool_name"] == "run_experiment" and not e["is_error"]]
        formats = [(format(r["test_accuracy"], ".4f"),
                    format(r["test_accuracy"] * 100, ".2f")) for r in results[-1]] if results else []
        checks["numerical_accuracy"] = bool(results) and all(
            re.search(rf"(?<!\d)(?:{re.escape(ratio)}|{re.escape(percent)}\s*%)(?!\d)", answer)
            is not None for ratio, percent in formats)
        numeric_only = dict(checks)
        if checks["numerical_accuracy"] != trial["checks"]["numerical_accuracy"]:
            notes.append("numeric-format-v2")
    if trial["case"].get("refusal"):
        checks["uncertainty_marker"] = bool(checks["uncertainty_marker"] or re.search(
            r"(?:无法|不能|难以)[^。！？\n]{0,100}(?:确认|确定|判断|给出|回答)", answer))
        if checks["uncertainty_marker"] != trial["checks"]["uncertainty_marker"]:
            notes.append("uncertainty-wording-v2")
    return checks, notes, all(numeric_only.values())


def percentile(values, q):
    values = [v for v in values if v is not None]
    return float(np.percentile(values, q)) if values else None


def summary_for(trials):
    calls = [c for t in trials for c in t["calls"]]
    turns = [u for t in trials for u in t["turns"]]
    starts = [e for t in turns for e in events(t, "ToolExecutionStarted")]
    completed = [e for t in turns for e in events(t, "ToolExecutionCompleted")]
    schema_valid = 0
    for e in starts:
        try:
            SCHEMAS[e["tool_name"]].model_validate(e["tool_input"])
            schema_valid += 1
        except (ValueError, KeyError):
            pass
    return {
        "n": len(trials), "raw_passed": sum(t["passed"] for t in trials),
        "numeric_only_passed": sum(t["numeric_only_passed"] for t in trials),
        "regraded_passed": sum(t["regraded_passed"] for t in trials),
        "p50_seconds": percentile([t["elapsed_seconds"] for t in trials], 50),
        "p95_seconds": percentile([t["elapsed_seconds"] for t in trials], 95),
        "p50_first_visible_seconds": percentile([u.get("first_visible_seconds") for u in turns], 50),
        "usage": usage_of(calls),
        "mean_observed_tokens": statistics.mean(t["meter"]["observed_total_tokens"] for t in trials),
        "tool_calls": len(starts), "schema_valid_calls": schema_valid,
        "tool_errors": sum(e["is_error"] for e in completed),
        "successful_tool_results": sum(not e["is_error"] for e in completed),
        "required_tools_passed": sum(t["checks"]["required_tools"] for t in trials),
        "message_pairs_passed": sum(u["messages_paired"] for u in turns), "turns": len(turns),
        "citation_gate_passed": sum(t["checks"]["citations_valid"] for t in trials),
        "max_request_chars": max(c["request_chars"] for c in calls),
        "compaction_attempts": sum(not c["request"].get("tools") for c in calls),
        "retrieval_calls": sum(e["tool_name"] in {"search_papers", "read_paper"} for e in starts),
        "context_facts_correct": sum(v for u in turns for v in u.get("fact_checks", {}).values()),
        "context_facts_total": sum(len(u.get("fact_checks", {})) for u in turns),
    }


def excel_value(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, float):
        return round(value, 4)
    if value is None:
        return "未知/不适用"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, str):
        # Prevent answer text from becoming spreadsheet formulas.
        value = value.replace("\x00", "")
        if value.startswith(("=", "+", "-", "@")):
            value = "'" + value
        if len(value) > 32000:
            return value[:31900] + "\n[完整文本见原始trace链接]"
    return value


def sheet(workbook, title, rows):
    ws = workbook.create_sheet(title)
    if not rows:
        rows = [{"说明": "无适用记录"}]
    ws.append(list(rows[0]))
    for row in rows:
        ws.append([excel_value(v) for v in row.values()])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 30
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="214F86")
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for number, header in enumerate(rows[0], 1):
        long = any(word in header for word in ("回答", "证据", "说明", "原因", "输入", "JSON", "原文"))
        ws.column_dimensions[get_column_letter(number)].width = 65 if long else 24
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="16314F")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EDF3FA")
            header = ws.cell(1, cell.column).value
            if isinstance(cell.value, (int, float)):
                if header == "通过比例":
                    cell.number_format = "0.0%"
                elif header.endswith("秒"):
                    cell.number_format = "0.00"
                elif "Token" in header:
                    cell.number_format = "#,##0"
            if header == "原始文件":
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
        longest = max(len(str(c.value)) for c in row)
        ws.row_dimensions[row[0].row].height = min(112, max(32, longest // 45 * 14))
    return ws


def plot_results(summary):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10,
                         "axes.spines.top": False, "axes.spines.right": False})
    fig, axes = plt.subplots(1, 3, figsize=(14, 4.5), layout="constrained")
    colors = ["#8EA4B7", "#4A7DAA", "#258378"]
    labels = ["V0", "V1", "V2"]
    for ax, group, title in zip(
        axes[:2], ("normal", "context"),
        ("Engineering checks (30 trials/version)", "Synthetic memory (12 trials/version)"),
        strict=True,
    ):
        rows = [summary[group][h] for h in NAMES]
        bars = ax.bar(labels, [r["regraded_passed"] / r["n"] * 100 for r in rows], color=colors)
        ax.bar_label(bars, labels=[f'{r["regraded_passed"]}/{r["n"]}' for r in rows], padding=5)
        ax.set(ylim=(0, 112), ylabel="Trials passing all checks (%)", title=title)
        ax.grid(axis="y", alpha=0.12)
        ax.set_axisbelow(True)
    data = [summary["context_by_case"]["tool_bloat"][h]["mean_observed_tokens"] for h in NAMES]
    bars = axes[2].bar(labels, np.array(data) / 1000, color=colors)
    axes[2].bar_label(bars, fmt="%.1fk", padding=5)
    axes[2].set(ylim=(0, 80), ylabel="Mean total tokens / 3-turn session (k)",
                title="Tool-result bloat (3 sessions/version)")
    fig.suptitle("Vision Scholar | Same model, prompts, tools and retrieval", fontsize=15)
    fig.text(.01, -.03, "Mechanical checks are not answer accuracy. Small synthetic cases; "
             "V2 evaluated after baseline. All tool-bloat usage complete.", fontsize=9)
    directory = OUTPUT / "figures"
    directory.mkdir(exist_ok=True)
    for suffix in ("png", "svg"):
        fig.savefig(directory / f"comparison.{suffix}", dpi=180, bbox_inches="tight")
    plt.close(fig)


def main():
    paths = sorted((OUTPUT / "trials").glob("*.json"))
    trials = [json.loads(p.read_text()) for p in paths]
    assert len([t for t in trials if t["group"] == "normal"]) == 90
    assert len([t for t in trials if t["group"] == "context"]) == 36
    store = Store(OUTPUT / "corpus.db")
    replays, task_rows, answer_rows, context_rows, call_rows, tool_rows, compact_rows = (
        [] for _ in range(7))
    for trial in trials:
        trial["regraded_checks"], trial["corrections"], trial["numeric_only_passed"] = regrade(trial)
        trial["regraded_passed"] = all(trial["regraded_checks"].values())
        trial["meter"] = usage_of(trial["calls"])
        ref = f"trials/{trial['id']}.json"
        base = {"样本ID": trial["id"], "组别": GROUPS[trial["group"]],
                "版本": NAMES[trial["harness"]], "任务": CASES[trial["case_id"]],
                "重复序号": trial["repetition"]}
        task_rows.append({**base, "原始通过": trial["passed"],
                          "数值校正后通过": trial["numeric_only_passed"],
                          "最终工程通过": trial["regraded_passed"],
                          "原始判据JSON": trial["checks"],
                          "修订判据JSON": trial["regraded_checks"], "修订原因": trial["corrections"],
                          "耗时秒": trial["elapsed_seconds"], **{
                              "总Token": trial["meter"]["total_tokens"],
                              "已观测Token小计": trial["meter"]["observed_total_tokens"],
                              "用量完整": trial["meter"]["usage_complete"]}, "原始文件": ref})
        for number, turn in enumerate(trial["turns"], 1):
            evidence = evidence_from(turn)
            answer_rows.append({**base, "轮次": number, "用户输入": turn["prompt"],
                                "完整回答": turn["answer"],
                                "引用有效ID": [c["id"] for c in turn["verification"]["valid"]],
                                "引用无效": turn["verification"]["invalid"],
                                "当前轮证据ID": list(evidence),
                                "首个可见文本秒": turn.get("first_visible_seconds"),
                                "本轮耗时秒": turn["elapsed_seconds"], "原始文件": ref})
            if trial["case_id"] == "code":
                after = verify_citations_with_ranges(turn["answer"], evidence, store)
                replays.append({
                    "trial_id": trial["id"], "group": trial["group"], "harness": trial["harness"],
                    "before": turn["verification"], "after": after,
                    "before_passed": (not turn["verification"]["invalid"]
                                      and bool(turn["verification"]["valid"])),
                    "after_passed": not after["invalid"] and bool(after["valid"]),
                    "basis": "固定同一回答与同一轮读取结果，只替换引用门禁；无模型重采样。",
                })
            if trial["group"] == "context":
                context_rows.append({**base, "轮次": number, "预设事实JSON": trial["case"]["facts"],
                                     "逐项事实判据": turn["fact_checks"],
                                     "完整回答": turn["answer"],
                                     "本轮耗时秒": turn["elapsed_seconds"],
                                     "首个可见文本秒": turn.get("first_visible_seconds"),
                                     "输出历史消息数": len(turn["history"]),
                                     "输出历史字符数": len(json.dumps(
                                         turn["history"], ensure_ascii=False)),
                                     "决策原文条数": len(turn.get("protected_decisions", [])),
                                     "原始文件": ref})
            pending = collections.defaultdict(list)
            for e in turn["events"]:
                if e["type"] == "ToolExecutionStarted":
                    pending[e["tool_name"]].append(e)
                elif e["type"] == "ToolExecutionCompleted":
                    before = pending[e["tool_name"]].pop(0)
                    tool_rows.append({**base, "轮次": number, "工具": e["tool_name"],
                                      "输入JSON": before["tool_input"], "报错": e["is_error"],
                                      "事件区间秒": e["seconds"] - before["seconds"],
                                      "结果字符数": len(e["output"]), "原始文件": ref})
        # The metadata dictionaries in early traces were shared across turn snapshots.
        # Read the final cumulative list once; never count it once per user turn.
        checkpoints = trial["turns"][-1].get("compact_metadata", {}).get("compact_checkpoints", [])
        for number, cp in enumerate(checkpoints, 1):
            compact_rows.append({**base, "累计序号": number, "分支": cp["checkpoint"],
                                 "元数据JSON": cp, "原始文件": ref})
        for number, call in enumerate(trial["calls"], 1):
            u = call["usage"] or {}
            call_rows.append({**base, "调用序号": number,
                              "调用类型": "模型执行" if call["request"].get("tools") else "上下文摘要",
                              "尝试序号": call["attempt"], "消息数": call["message_count"],
                              "请求字符数": call["request_chars"], "用量已返回": bool(u),
                              "输入Token": u.get("input_tokens"), "输出Token": u.get("output_tokens"),
                              "总Token": u.get("total_tokens"), "缓存Token": u.get("cached_tokens"),
                              "推理Token": u.get("reasoning_tokens"), "错误": call.get("error", ""),
                              "耗时秒": call.get("elapsed_seconds"), "原始文件": ref})
    summary = {}
    for group in ("normal", "context", "canary"):
        summary[group] = {
            h: summary_for([t for t in trials if t["group"] == group and t["harness"] == h])
            for h in NAMES}
    summary["context_by_case"] = {
        case: {h: summary_for([t for t in trials if t["group"] == "context"
                              and t["case_id"] == case and t["harness"] == h]) for h in NAMES}
        for case in ("short", "tool_bloat", "buried_correction", "full_summary")}
    summary["normal_by_case"] = {
        case: {h: summary_for([t for t in trials if t["group"] == "normal"
                              and t["case_id"] == case and t["harness"] == h]) for h in NAMES}
        for case in CASES if any(t["group"] == "normal" and t["case_id"] == case for t in trials)}
    review_records = [json.loads(p.read_text())
                      for p in (OUTPUT / "semantic-review").glob("normal-*.json")]
    assert len(review_records) == 30 and all(r.get("review") for r in review_records)
    review_rows = []
    summary["semantic_review"] = {}
    for h in NAMES:
        selected = [r for r in review_records if f"-{h}-" in r["trial_id"]]
        summary["semantic_review"][h] = {
            "n": len(selected), "verdicts": dict(collections.Counter(
                r["review"]["verdict"] for r in selected)),
            "basis": "首次重复分层选样；匿名同模型AI复核；非人评金标或准确率。"}
    for record in review_records:
        review_rows.append({
            "样本ID": record["trial_id"], "事实支持0至4": record["review"]["factual_support"],
            "指令完成0至4": record["review"]["instruction_completion"],
            "边界诚实0至4": record["review"]["boundary_honesty"],
            "AI判定": record["review"]["verdict"], "AI判断理由": record["review"]["reason"],
            "AI问题及证据JSON": record["review"]["issues"],
            "原始文件": f"semantic-review/{record['trial_id']}.json"})
    reviews_usage = [a["usage"] for r in review_records for a in r["attempts"] if a.get("usage")]
    summary["review_usage"] = {
        "attempts": sum(len(r["attempts"]) for r in review_records),
        "usage_records": len(reviews_usage),
        "observed_total_tokens": sum(u["total_tokens"] for u in reviews_usage)}
    audit = json.loads((OUTPUT / "agent-evidence-audit.json").read_text())
    summary["agent_audit"] = audit
    summary["corrections"] = CORRECTIONS
    summary["code_gate_replay"] = [
        {k: r[k] for k in ("trial_id", "harness", "before_passed", "after_passed")}
        for r in replays if r["group"] == "normal"]
    summary["normalizations"] = {
        "usage": "按calls[].usage重算；缺失则总量未知，另列已观测小计。",
        "ttft": "用户轮开始至AssistantTextDelta；不使用含隐藏摘要的transport首文本。",
        "compact": "早期turn元数据引用同一可变对象，只计最后一轮的累计列表一次。",
        "tool_time": "同名并发调用事件无call id，按同名开始顺序配对；"
                     "上游多工具完成事件在整批收齐后发送，事件区间不是单工具纯执行时间。",
        "sample": "10固定正常任务×3重复×3版本；4合成上下文×3重复×3版本；9连通小样另列。",
    }
    summary["trials"] = [{
        "id": t["id"], "group": t["group"], "harness": t["harness"], "case_id": t["case_id"],
        "raw_checks": t["checks"], "regraded_checks": t["regraded_checks"],
        "regraded_passed": t["regraded_passed"], "meter": t["meter"],
    } for t in trials]
    write_json(OUTPUT / "comparison-summary.json", summary)
    write_json(OUTPUT / "citation-gate-ablation.json", replays)
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "阅读说明", [
        {"项目": "实验对象", "说明": "V0按原简历重建独立LangGraph循环；V1复用OpenHarness；"
         "V2增加用户决策原文、代码子范围验证、单run笔记去重。"},
        {"项目": "可比较范围", "说明": summary["normalizations"]["sample"]},
        {"项目": "控制变量", "说明": "doubao-seed-2-0-lite-260428；temperature=0；thinking关闭；"
         "相同系统Prompt、工具、论文、预热LSA；V2基于基线发现后单独批次运行。"},
        {"项目": "评测依据", "说明": "程序工程判据与AI语义复核分别报告；没有人评金标。"},
        {"项目": "用量", "说明": summary["normalizations"]["usage"]},
        {"项目": "时间", "说明": summary["normalizations"]["ttft"]},
        {"项目": "压缩", "说明": summary["normalizations"]["compact"]},
        {"项目": "检索", "说明": "48条测量=8固定英文问题×3重复×2检索器；限定论文和预设页。"
         "重复不是24个独立问题，神经检索单独比较。"},
        {"项目": "故障", "说明": "脚本Mock和人工延迟测工程机制，不算真实模型能力。"},
        {"项目": "边界", "说明": "当前仍为单用户单worker、受限源码、固定CPU实验；"
         "没有自动中断工具续跑，决策关键词规则不能保证所有自然语言修正被识别。"},
    ])
    overview = []
    for group in ("normal", "context"):
        for h, s in summary[group].items():
            overview.append({
                "组别": GROUPS[group], "版本": NAMES[h], "样本数": s["n"],
                "原始通过数": s["raw_passed"], "最终工程通过数": s["regraded_passed"],
                "通过比例": s["regraded_passed"] / s["n"], "耗时中位秒": s["p50_seconds"],
                "耗时P95秒": s["p95_seconds"], "首文本中位秒": s["p50_first_visible_seconds"],
                "完整总Token": s["usage"]["total_tokens"],
                "已观测Token小计": s["usage"]["observed_total_tokens"],
                "用量完整": s["usage"]["usage_complete"], "工具调用数": s["tool_calls"],
                "参数合法数": s["schema_valid_calls"], "工具结果错误数": s["tool_errors"],
                "必要工具完成任务数": s["required_tools_passed"],
                "消息配对通过轮数": s["message_pairs_passed"],
                "引用门禁通过任务数": s["citation_gate_passed"],
                "摘要API尝试数": s["compaction_attempts"],
            })
    sheet(wb, "总览", overview)
    sheet(wb, "分任务汇总", [{
        "组别": GROUPS[group], "任务": CASES[case], "版本": NAMES[h],
        "重复数": r["n"], "工程通过数": r["regraded_passed"],
        "耗时中位秒": r["p50_seconds"], "平均已观测Token": r["mean_observed_tokens"],
        "用量完整": r["usage"]["usage_complete"], "事实正确项数": r["context_facts_correct"],
        "事实总项数": r["context_facts_total"],
    } for group in ("normal", "context") for case, rows in summary[group + "_by_case"].items()
       for h, r in rows.items()])
    sheet(wb, "任务明细", task_rows)
    sheet(wb, "回答与证据", answer_rows)
    sheet(wb, "上下文逐轮", context_rows)
    sheet(wb, "调用用量", call_rows)
    sheet(wb, "工具执行", tool_rows)
    sheet(wb, "压缩分支", compact_rows)
    sheet(wb, "代码门禁消融", [{
        "样本ID": r["trial_id"], "版本": NAMES[r["harness"]], "原门禁通过": r["before_passed"],
        "新门禁通过": r["after_passed"], "新门禁证据JSON": r["after"],
        "说明": r["basis"], "原始文件": "citation-gate-ablation.json",
    } for r in replays])
    retrieval = json.loads((OUTPUT / "retrieval.json").read_text())
    sheet(wb, "检索对照", [{
        "检索方案": r["retriever"], "问题ID": r["id"], "重复序号": r["repetition"],
        "输入查询": r["query"], "限定论文": r["paper_id"], "预设相关页": r["reference_pages"],
        "首个相关排名": r["first_relevant_rank"], "前六命中": r["hit_at_6"],
        "倒数排名": r["rr_at_6"], "耗时秒": r["elapsed_seconds"], "原始文件": "retrieval.json",
    } for r in retrieval["rows"]])
    faults = json.loads((OUTPUT / "faults.json").read_text())
    sheet(wb, "故障回归", [{"测试": r["name"], "通过": r["passed"],
                         "耗时秒": r["seconds"], "原始文件": "faults.json"} for r in faults["tests"]])
    sheet(wb, "模拟并行", [{
        "版本": NAMES[r["harness"]], "重复序号": r["repetition"], "总耗时秒": r["total_seconds"],
        "三工具耗时区间秒": r["tool_span_seconds"], "消息配对": r["messages_paired"],
        "说明": faults["basis"], "原始文件": "faults.json",
    } for r in faults["scheduling"]])
    sheet(wb, "AI语义复核", review_rows)
    sheet(wb, "定向证据回查", [{
        "样本ID": r["trial_id"], "问题程度": r["severity"], "回答原文": r["answer_quote"],
        "证据ID": r["evidence_id"], "证据原文": r["evidence_quote"],
        "回查判断": r["assessment"], "原匿名AI判定": r["blind_review_verdict"],
        "原始文件": r["trace"],
    } for r in audit["findings"]])
    sheet(wb, "判据修订", [{"修订ID": c["id"], "原因": c["reason"], "范围": c["scope"]}
                          for c in CORRECTIONS])
    wb.save(OUTPUT / "Harness对比与证据.xlsx")
    plot_results(summary)
    for group in ("normal", "context"):
        for h, s in summary[group].items():
            print(group, h, "pass", s["regraded_passed"], "/", s["n"],
                  "p50", round(s["p50_seconds"], 2), "observed",
                  s["usage"]["observed_total_tokens"], "complete", s["usage"]["usage_complete"])
    print("AI review:", summary["semantic_review"])
    print("Artifacts:", OUTPUT / "Harness对比与证据.xlsx")


if __name__ == "__main__":
    main()
